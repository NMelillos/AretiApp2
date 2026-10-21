"""Report-only date provenance and protected baseline parity; disposable data."""
from contextlib import closing
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import ast,importlib.util,json,os
from pathlib import Path
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
import db
from _qa_transaction_import_date import DateTests
from _qa_balance_date_validation import DatePresentationTests,VisibilityTests,BalanceDatePreservationTests
from _qa_reports_five_placement import ReportsPlacementTests
from _qa_autonomous_workflows import BalanceActionTests,CompactTests
from _qa_autonomous_zero_boc import ZeroSourceTests
from _qa_balance_cnb_report import BalanceTests
from _qa_balance_print_ungrouped import PrintTests
from latest_import_balances import snapshot,print_document,workbook_bytes,COLUMNS,_balance_import_date
from latest_balances_compact import compact_html,compact_model
from import_history import cyprus_time

BASELINE=Path(__file__).resolve().parent.parent/'areti-balances-finalisation'


class AssociationTests(DateTests):
    def test_normal_import_has_three_distinct_date_sources(self):
        self.seed('normal','2026-10-01T08:00:00+00:00',True)
        row=self.row()
        self.assertEqual(row['Transaction Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))
        self.assertEqual(row['Balance Import Date'],row['Transaction Import Date'])
        self.assertEqual(row['Statement end date'],'2026-09-30')

    def test_later_balance_record_does_not_replace_transaction_timestamp(self):
        self.seed('transactions','2026-10-01T08:00:00+00:00',True)
        self.seed('balance','2026-10-15T08:00:00+00:00')
        row=self.row()
        self.assertEqual(row['Transaction Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))
        self.assertEqual(row['Balance Import Date'],cyprus_time('2026-10-15T08:00:00+00:00'))

    def test_verified_overlap_uses_its_referenced_transaction_history(self):
        from balance_only import PREFIX
        self.seed('original','2026-10-01T08:00:00+00:00',True)
        self.seed('unrelated','2026-10-03T08:00:00+00:00',True)
        with closing(db.get_connection()) as c:
            ident=c.execute("SELECT id FROM classified_transactions WHERE statement_hash='original'").fetchone()[0]
        notes=PREFIX+json.dumps({'references':[[ident,'2026-09-15','-1.00']]})
        self.seed('overlap','2026-10-15T08:00:00+00:00',opening='10.00',closing_balance='9.00',notes=notes)
        with closing(db.get_connection()) as c:c.execute("UPDATE statement_balances SET money_out='1.00' WHERE statement_hash='overlap'");c.commit()
        row=self.row();self.assertEqual(row['Status'],'IMPORTED')
        self.assertEqual(row['Transaction Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))
        self.assertEqual(row['Balance Import Date'],cyprus_time('2026-10-15T08:00:00+00:00'))

    def test_null_transaction_count_never_crashes_or_fabricates_event(self):
        self.seed('legacy','2026-10-01T08:00:00+00:00',True)
        with closing(db.get_connection()) as c:c.execute('UPDATE statement_imports SET transaction_count=NULL');c.commit()
        row=self.row();self.assertEqual(row['Status'],'INCOMPLETE')
        self.assertEqual(row['Transaction Import Date'],'—')

    def test_new_zero_statement_preserves_old_transaction_event_without_warning(self):
        self.seed('transactions','2026-07-01T08:00:00+00:00',True)
        self.seed('zero','2026-10-15T08:00:00+00:00',opening='25.00')
        rows,total,warnings=self.report();row=rows[0]
        self.assertEqual(row['Transaction Import Date'],cyprus_time('2026-07-01T08:00:00+00:00'))
        self.assertEqual(row['Balance Import Date'],cyprus_time('2026-10-15T08:00:00+00:00'))
        self.assertEqual(row['Closing balance'],Decimal('25.00'));self.assertEqual(warnings,[])

    def older_import(self):
        self.seed('newer-period','2026-10-01T08:00:00+00:00',True)
        self.seed('older-period','2026-10-15T08:00:00+00:00',True)
        with closing(db.get_connection()) as c:
            c.execute("UPDATE statement_balances SET source='BOC bank columns',bank='Bank of Cyprus'")
            c.execute("UPDATE account_list SET bank='Bank of Cyprus'")
            c.execute("UPDATE classified_transactions SET bank='Bank of Cyprus'")
            c.execute("UPDATE statement_balances SET period_start='2026-08-01',period_end='2026-08-31' WHERE statement_hash='older-period'")
            c.execute("UPDATE classified_transactions SET txn_date='2026-08-15' WHERE statement_hash='older-period'")
            c.commit()

    def test_selected_statement_owns_all_dates_when_older_import_arrives_later(self):
        self.older_import();row=self.row()
        self.assertEqual(row['Statement end date'],'2026-09-30')
        self.assertEqual(row['Transaction Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))
        self.assertEqual(row['Balance Import Date'],row['Transaction Import Date'])

    def test_missing_balance_timestamp_never_borrows_transaction_or_update(self):
        self.seed('normal','2026-10-01T08:00:00+00:00',True)
        with closing(db.get_connection()) as c:
            c.execute('UPDATE statement_balances SET imported_at=NULL');c.commit()
        row=self.row();self.assertEqual(row['Balance Import Date'],'—')
        self.assertTrue(row['Transaction Import Date'].startswith('2026-10-01'))

    def test_generic_update_is_not_a_balance_import(self):
        self.seed('normal','2026-10-01T08:00:00+00:00',True)
        with closing(db.get_connection()) as c:
            c.execute("UPDATE statement_balances SET updated_at='2026-10-20T08:00:00+00:00'");c.commit()
        row=self.row();self.assertEqual(row['Balance Import Date'],'—')
        self.assertEqual(row['Transaction Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))
        self.assertTrue(any('later metadata update' in w for w in self.report()[2]))

    def test_source_bound_completion_audit_and_wrong_hash_rejection(self):
        self.seed('normal','2026-10-01T08:00:00+00:00',True)
        audit=dict(action='user-confirmed BOC balances only',actor='Synthetic reviewer',source_hash='normal',transaction_changes=0,confirmed_at='2026-10-15T08:00:00+00:00')
        with closing(db.get_connection()) as c:
            c.execute("UPDATE statement_balances SET source='BOC bank columns',notes=?,updated_at=?",('Existing preserved note\n'+json.dumps(audit),audit['confirmed_at']));c.commit()
        self.assertEqual(self.row()['Balance Import Date'],cyprus_time(audit['confirmed_at']))
        audit['source_hash']='other-statement'
        with closing(db.get_connection()) as c:c.execute('UPDATE statement_balances SET notes=?',(json.dumps(audit),));c.commit()
        self.assertEqual(self.row()['Balance Import Date'],'—')

    def test_invalid_or_outdated_audit_cannot_supply_date(self):
        base=dict(balance_statement_hash='normal',balance_imported_at='2026-10-01T08:00:00+00:00',balance_updated_at='2026-10-15T08:00:00+00:00',source='BOC bank columns')
        for event in (None,[],{'action':'other'},{'action':'user-confirmed BOC balances only','source_hash':'normal','actor':'Reviewer','transaction_changes':0,'confirmed_at':'not-a-date'}, {'action':'user-confirmed BOC balances only','source_hash':'normal','actor':'Reviewer','transaction_changes':0,'confirmed_at':'2026-10-10T08:00:00+00:00'}):
            row=dict(base,notes=json.dumps(event))
            self.assertEqual(_balance_import_date(row)[0],'—')

    def test_legacy_timezone_is_explicit_and_invalid_date_is_unavailable(self):
        for stamp,expected in (('2026-10-01 08:00:00','2026-10-01 08:00:00 (timezone unverified)'),('bad-date','—'),(None,'—')):
            row=dict(balance_statement_hash='normal',balance_imported_at=stamp,balance_updated_at=stamp,source='Synthetic summary',notes='')
            self.assertEqual(_balance_import_date(row)[0],expected)

    def test_bank_and_currency_do_not_supply_other_accounts_balance_date(self):
        self.seed('own','2026-10-01T08:00:00+00:00',True)
        self.seed('other-bank','2026-10-20T08:00:00+00:00',bank='Bank B')
        self.seed('other-currency','2026-10-21T08:00:00+00:00',currency='EUR')
        self.assertEqual(self.row()['Balance Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))

    def test_ui_print_excel_share_balance_date_and_preserve_formatting(self):
        self.seed('normal','2026-10-01T08:00:00+00:00',True)
        rows,total,warnings=self.report();before=repr(rows)
        now=datetime.fromisoformat('2026-10-21T12:00:00+03:00')
        html=compact_html(rows,now);doc=print_document(rows,total,warnings,now)
        labels=('Bank','Account','Account Number','Transaction Import Date','Balance Import Date','Statement Closing Date','Currency','Closing Balance','Closing Balance USD')
        positions=[html.index('>'+label+'</th>') for label in labels]
        self.assertEqual(positions,sorted(positions));self.assertIn('colspan="8"',html)
        self.assertIn(rows[0]['Balance Import Date'],html);self.assertIn('<th>Balance Import Date</th>',doc)
        self.assertNotIn('class="group"',doc);self.assertIn('width:max-content',html);self.assertIn('9.00',html)
        book=load_workbook(BytesIO(workbook_bytes(rows,total,warnings,now)));col=COLUMNS.index('Balance Import Date')+1
        self.assertEqual(book.active.cell(2,col).value,rows[0]['Balance Import Date'])
        self.assertEqual(book.active.cell(2,col).font.color.rgb,'00000000')
        self.assertEqual(repr(rows),before)

    def test_sort_total_and_financial_fields_equal_production_snapshot(self):
        self.older_import()
        spec=importlib.util.spec_from_file_location('pinned_balance_report',BASELINE/'latest_import_balances.py')
        pinned=importlib.util.module_from_spec(spec);spec.loader.exec_module(pinned)
        before=self.dump();old,old_total,old_warn=pinned.snapshot(db);new,new_total,new_warn=snapshot(db)
        self.assertEqual(new_total,old_total)
        for a,b in zip(old,new):
            self.assertEqual({k:v for k,v in a.items() if k!='Transaction Import Date'},{k:v for k,v in b.items() if k not in ('Transaction Import Date','Balance Import Date')})
        self.assertEqual(self.dump(),before)
        self.assertEqual([w for w in new_warn if 'Balance Import Date unavailable:' not in w],old_warn)

    def test_overlap_history_query_is_bound_readonly_and_missing_history_not_borrowed(self):
        from latest_import_balances import _referenced_transaction_events
        from balance_only import PREFIX
        from unittest.mock import Mock
        from pglast import parse_sql
        conn=Mock();cur=conn.cursor.return_value
        row=dict(notes=PREFIX+json.dumps({'references':[[12,'2026-09-01','-1.00']]}),statement_bank='Bank A',statement_account_number='A1',statement_currency='USD')
        cur.fetchall.return_value=[(None,None,None,1)]
        self.assertEqual(_referenced_transaction_events(conn,row),[])
        sql,args=cur.execute.call_args.args
        self.assertEqual(args,(12,'Bank A','A1','USD'))
        parsed=parse_sql(sql.replace('?','NULL'))
        self.assertEqual(len(parsed),1);self.assertEqual(type(parsed[0].stmt).__name__,'SelectStmt')
        cur.close.assert_called_once()


class ActualCompletionTests(BalanceDatePreservationTests):
    def test_balance_completion_audit_sets_date_without_rewriting_history_or_repeat(self):
        from boc_summary_balances import record
        path,content,frame,balance,digest=self.original()
        with closing(db.get_connection()) as c:
            c.execute("UPDATE statement_imports SET imported_at='2026-07-01T08:00:00+00:00' WHERE statement_hash=?",(digest,))
            c.execute("UPDATE statement_balances SET imported_at='2026-07-01T08:00:00+00:00',updated_at='2026-07-01T08:00:00+00:00' WHERE statement_hash=?",(digest,))
            c.execute("UPDATE classified_transactions SET reviewed=1,category='Synthetic',subcategory='General' WHERE statement_hash=?",(digest,));c.commit()
        before=self.unchanged_transactions()
        self.assertEqual(record(db,content,path.name,'Synthetic reviewer'),'recorded')
        row=next(r for r in snapshot(db)[0] if r['Account number']==balance['account_number'])
        self.assertTrue(row['Transaction Import Date'].startswith('2026-07-01'))
        self.assertNotEqual(row['Balance Import Date'],'—');self.assertNotEqual(row['Balance Import Date'],row['Transaction Import Date'])
        self.assertEqual(self.unchanged_transactions(),before)
        saved=self.snapshot();self.assertEqual(record(db,content,path.name,'Synthetic reviewer'),'already recorded')
        self.assertEqual(self.snapshot(),saved)
        self.assertEqual(next(r for r in snapshot(db)[0] if r['Account number']==balance['account_number'])['Balance Import Date'],row['Balance Import Date'])


class SourceScopeTests(unittest.TestCase):
    def test_only_report_layer_application_modules_changed(self):
        repo=Path(__file__).resolve().parent
        for file in BASELINE.glob('*.py'):
            if file.name in ('latest_import_balances.py','latest_balances_compact.py') or file.name.startswith('_qa'):continue
            self.assertEqual((repo/file.name).read_bytes(),file.read_bytes(),file.name)
        self.assertEqual((repo/'requirements.txt').read_bytes(),(BASELINE/'requirements.txt').read_bytes())
        before=ast.parse((BASELINE/'latest_import_balances.py').read_text(encoding='utf-8-sig'))
        after=ast.parse((repo/'latest_import_balances.py').read_text(encoding='utf-8-sig'))
        unchanged=('fresh_closing','fresh_import','_verified_time','_account_identity','_report_bank_equal','_sql_bank','_sql_number')
        for name in unchanged:
            a=next(n for n in before.body if isinstance(n,ast.FunctionDef) and n.name==name)
            b=next(n for n in after.body if isinstance(n,ast.FunctionDef) and n.name==name)
            self.assertEqual(ast.dump(a),ast.dump(b),name)


def suite():
    tests=unittest.TestSuite()
    tests.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(DateTests))
    for name in AssociationTests.__dict__:
        if name.startswith('test_'):tests.addTest(AssociationTests(name))
    tests.addTest(ActualCompletionTests('test_balance_completion_audit_sets_date_without_rewriting_history_or_repeat'))
    tests.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(SourceScopeTests))
    for cls,names in (
        (DatePresentationTests,('test_old_transaction_before_closing_is_not_a_date_anomaly','test_semantics_available_in_ui_print_and_excel_without_inferred_activity','test_closing_freshness_stays_independent_in_every_export')),
        (VisibilityTests,('test_third_hidden_even_with_stale_open_session','test_areti_standalone_view_still_available_and_readonly')),
        (ReportsPlacementTests,('test_authenticated_reports_button_compact_return_read_only','test_compact_error_keeps_return_and_no_write','test_reports_requires_authentication_and_preserves_login_route')),
        (BalanceActionTests,('test_conflicts_missing_rows_actor_and_rollback_fail_safely',)),
        (CompactTests,('test_visible_usd_sum_missing_fx_exclusion_no_wrap_or_warning_wall',)),
        (ZeroSourceTests,('test_real_synthetic_zero_pdf_source_proof_commit_and_repeat',)),
        (BalanceTests,('test_verified_overlap_atomic_no_new_rows_history_report_duplicate','test_reference_mutation_invalidates_balance_only_record','test_older_zero_row_statement_does_not_displace_newer')),
    ):tests.addTests(cls(name) for name in names)
    tests.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(PrintTests))
    return tests


if __name__=='__main__':
    result=unittest.TextTestRunner(verbosity=2).run(suite())
    Path(os.environ['ARETI_ISOLATION_ROOT'],'test-counts.json').write_text(json.dumps(dict(tests=result.testsRun,failures=len(result.failures),errors=len(result.errors),skipped=len(result.skipped))))
    raise SystemExit(not result.wasSuccessful())
