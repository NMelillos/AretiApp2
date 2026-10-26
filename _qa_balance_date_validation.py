"""Presentation-only package; credential-free disposable SQLite tests."""
import ast
from contextlib import closing
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
import os
from pathlib import Path
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
import db
from _qa_transaction_import_date import DateTests
from _qa_post_import_message import CompletionTests
from _qa_reports_five_placement import ReportsPlacementTests
from _qa_autonomous_workflows import BalanceActionTests
from latest_import_balances import NOTE, snapshot, print_document, workbook_bytes, fresh_closing
from latest_balances_compact import compact_html


class DatePresentationTests(DateTests):
    def test_old_transaction_before_closing_is_not_a_date_anomaly(self):
        self.seed('old-transactions','2026-07-01T08:00:00+00:00',True)
        self.seed('new-balances','2026-10-10T08:00:00+00:00')
        rows,total,warnings=self.report()
        row=rows[0]
        self.assertTrue(row['Transaction Import Date'].startswith('2026-07-01'))
        self.assertEqual(row['Statement end date'],'2026-09-30')
        self.assertEqual(warnings,[])
        generated=datetime.fromisoformat('2026-10-17T12:00:00+03:00')
        html=print_document(rows,total,warnings,generated)
        self.assertIn('style="color:#000000">'+row['Transaction Import Date'],html)
        self.assertIn('color:#000000',compact_html(rows,generated))
        book=load_workbook(BytesIO(workbook_bytes(rows,total,warnings,generated)))
        self.assertEqual(book.active.cell(2,4).font.color.rgb,'00000000')
        self.assertEqual(book.active.cell(2,6).font.color.rgb,'00146B36')
        self.assertEqual(total,Decimal('10.00'))

    def test_semantics_available_in_ui_print_and_excel_without_inferred_activity(self):
        self.seed('balance-only','2026-10-10T08:00:00+00:00')
        rows,total,warnings=self.report()
        self.assertEqual(rows[0]['Transaction Import Date'],'—')
        self.assertIn('zero history count alone does not prove zero source activity',NOTE)
        for label in ('Statement Closing Date','Transaction Import Date','Balance Import Date','General Report Date'):
            self.assertIn(label,NOTE)
        self.assertIn(NOTE,print_document(rows,total,warnings))
        book=load_workbook(BytesIO(workbook_bytes(rows,total,warnings)))
        self.assertIn(NOTE,[c.value for row in book['Verification'] for c in row])

    def test_closing_freshness_stays_independent_in_every_export(self):
        self.seed('old-transactions','2026-07-01T08:00:00+00:00',True)
        rows,total,warnings=self.report()
        now=datetime.fromisoformat('2026-10-17T12:00:00+03:00')
        for age in (-1,0,29,30,31,40):
            row=dict(rows[0])
            closed=(now.date()-timedelta(days=age)).isoformat()
            row['Statement end date']=closed
            fresh=0<=age<=30
            self.assertEqual(fresh_closing(closed,now),fresh)
            self.assertIn('class="'+('fresh' if fresh else 'stale')+'">'+closed,print_document([row],total,warnings,now))
            html=compact_html([row],now)
            self.assertIn('color:'+('#146b36' if fresh else '#b42318'),html)
            book=load_workbook(BytesIO(workbook_bytes([row],total,warnings,now)))
            self.assertEqual(book.active.cell(2,4).font.color.rgb,'00000000')
            self.assertEqual(book.active.cell(2,6).font.color.rgb,'00146B36' if fresh else '00B42318')
        for value in ('',None,'invalid'):
            self.assertFalse(fresh_closing(value,now))


class VisibilityTests(CompletionTests):
    def test_third_hidden_even_with_stale_open_session(self):
        before=self.snapshot()
        app=AppTest.from_file('app.py',default_timeout=45)
        app.query_params['page']='TB & NF Family Office Report'
        app.session_state['third_report_authenticated']=True
        app.session_state['third_report_user']='Synthetic report operator'
        app.session_state['third_latest_balances_open']=True
        app.run()
        self.assertFalse(app.exception,[e.message for e in app.exception])
        self.assertFalse(any(b.key in ('third_latest_import_balances','third_latest_balances_return') for b in app.button))
        self.assertFalse(any('5. Latest Import Balances' in m.value for m in app.markdown))
        self.assertTrue(any('No transactions are available' in m.value for m in app.info))
        self.assertEqual(self.snapshot(),before)

    def test_areti_standalone_view_still_available_and_readonly(self):
        before=self.snapshot()
        app=AppTest.from_file('app.py',default_timeout=45)
        app.query_params['page']='Latest Import Balances'
        app.session_state['authenticated']=True
        app.session_state['login_user']='Synthetic Areti reviewer'
        with patch.object(db,'backfill_missing_usd_amounts',side_effect=AssertionError('Report must not backfill')):
            app.run()
        self.assertFalse(app.exception,[e.message for e in app.exception])
        self.assertIn('Latest Import Balances',[s.value for s in app.subheader])
        from html import unescape
        self.assertIn(NOTE,unescape(''.join(s.value for s in app.markdown)))
        self.assertTrue(any(e.label=='Report scope, date guidance and verification details' for e in app.expander))
        self.assertEqual(self.snapshot(),before)


class ScopeTests(unittest.TestCase):
    def test_only_third_visibility_removed_from_app(self):
        baseline=ast.parse((Path(os.environ['ARETI_DATE_BASELINE'])/'app.py').read_text(encoding='utf-8-sig'))
        current=ast.parse(Path('app.py').read_text(encoding='utf-8-sig'))
        baseline.body=[n for n in baseline.body if not isinstance(n,ast.FunctionDef) or n.name!='_render_third_balance_button']
        third=next(n for n in baseline.body if isinstance(n,ast.FunctionDef) and n.name=='render_third_link_report')
        class RemoveVisibility(ast.NodeTransformer):
            def visit_If(self,node):
                if ast.unparse(node.test)=="st.session_state.get('third_latest_balances_open')":return None
                return self.generic_visit(node)
            def visit_Expr(self,node):
                if isinstance(node.value,ast.Call) and ast.unparse(node.value.func)=='_render_third_balance_button':return None
                return node
        RemoveVisibility().visit(third)
        self.assertEqual(ast.dump(baseline),ast.dump(current))
        reports=lambda tree:next(n for n in ast.walk(tree) if isinstance(n,ast.If) and ast.unparse(n.test)=="page == 'Reports'")
        self.assertEqual(ast.dump(reports(baseline)),ast.dump(reports(current)))

    def test_report_logic_and_all_other_modules_identical_to_live(self):
        reference=Path(os.environ['ARETI_DATE_BASELINE'])
        for path in reference.glob('*.py'):
            if path.name.startswith('_qa') or path.name in ('app.py','latest_import_balances.py'):continue
            # Git checks this new Windows worktree out with CRLF; compare source
            # with universal newlines, preserving every other character.
            self.assertEqual(path.read_text(encoding='utf-8-sig'),Path(path.name).read_text(encoding='utf-8-sig'),path.name)
        def without_note(path):
            tree=ast.parse(path.read_text(encoding='utf-8-sig'))
            tree.body=[n for n in tree.body if not isinstance(n,ast.Assign) or not any(isinstance(t,ast.Name) and t.id=='NOTE' for t in n.targets)]
            return ast.dump(tree)
        self.assertEqual(without_note(reference/'latest_import_balances.py'),without_note(Path('latest_import_balances.py')))
        self.assertEqual(Path(db.__file__).resolve(),Path('db.py').resolve())


class BalanceDatePreservationTests(BalanceActionTests):
    def test_actual_balance_completion_keeps_classifications_and_transaction_date(self):
        from boc_summary_balances import record
        path,content,frame,balance,digest=self.original()
        with closing(db.get_connection()) as conn:
            conn.execute("UPDATE statement_imports SET imported_at='2026-07-01T08:00:00+00:00' WHERE statement_hash=?",(digest,))
            conn.execute("UPDATE classified_transactions SET reviewed=1, category='Synthetic', subcategory='General' WHERE statement_hash=?",(digest,))
            conn.commit()
        before=self.unchanged_transactions()
        date_before=next(r for r in snapshot(db)[0] if r['Account number']==balance['account_number'])['Transaction Import Date']
        self.assertTrue(date_before.startswith('2026-07-01'))
        self.assertEqual(record(db,content,path.name,'Synthetic reviewer'),'recorded')
        self.assertEqual(self.unchanged_transactions(),before)
        row=next(r for r in snapshot(db)[0] if r['Account number']==balance['account_number'])
        self.assertEqual(row['Transaction Import Date'],date_before)
        self.assertEqual(Decimal(str(row['Closing balance'])),balance['closing_balance'])
        self.assertEqual(row['Statement end date'],balance['period_end'])
        saved=self.snapshot()
        self.assertEqual(record(db,content,path.name,'Synthetic reviewer'),'already recorded')
        self.assertEqual(self.snapshot(),saved)


def suite():
    tests=unittest.TestSuite()
    for cls,names in (
        (DatePresentationTests,('test_old_transaction_before_closing_is_not_a_date_anomaly','test_semantics_available_in_ui_print_and_excel_without_inferred_activity','test_closing_freshness_stays_independent_in_every_export')),
        (VisibilityTests,('test_third_hidden_even_with_stale_open_session','test_areti_standalone_view_still_available_and_readonly')),
        (ScopeTests,('test_only_third_visibility_removed_from_app','test_report_logic_and_all_other_modules_identical_to_live')),
        (BalanceDatePreservationTests,('test_actual_balance_completion_keeps_classifications_and_transaction_date',)),
        (ReportsPlacementTests,('test_authenticated_reports_button_compact_return_read_only','test_compact_error_keeps_return_and_no_write','test_reports_requires_authentication_and_preserves_login_route')),
    ):
        tests.addTests(cls(name) for name in names)
    tests.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(DateTests))
    return tests


if __name__=='__main__':
    raise SystemExit(not unittest.TextTestRunner(verbosity=2).run(suite()).wasSuccessful())
