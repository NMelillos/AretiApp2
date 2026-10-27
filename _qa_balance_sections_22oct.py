"""Exact approved buckets; preserve source values, FX, selection and report dates."""
import importlib.util,os,unittest
from pathlib import Path
from decimal import Decimal
from io import BytesIO
from contextlib import closing
from datetime import datetime
from openpyxl import load_workbook
import db
from _qa_transaction_import_date import DateTests
from balance_report_sections import section,totals,partition
from latest_import_balances import snapshot,print_document,workbook_bytes,COLUMNS
from latest_balances_compact import compact_html

class SectionTests(DateTests):
    def model(self,bank,number,value='10.00',verification='SOURCE RECONCILIATION NOT VERIFIED'):
        row={c:'' for c in COLUMNS}
        row.update({'Bank':bank,'Account name':'Synthetic','Account number':number,'Currency':'USD','Closing balance':Decimal(value),'Closing balance converted to USD':Decimal(value),'Status':'IMPORTED','Verification':verification,'Statement end date':'2026-09-30','Transaction Import Date':'2026-10-01','Balance Import Date':'2026-10-22'})
        return row

    def test_exact_debt_whitelist_and_no_all_comerica_or_citi_rule(self):
        debts=[('AMEX','A'),('American Express','A'),('Citi Bank Credit Card','B'),('Sapphire Chase','C'),('Comerica','055-9000-0991 Credit line')]
        deposits=[('Comerica','other'),('Comerica','055-9000-0992 Credit line'),('CNB','D'),('Revolut','E'),('Citi','F'),('Fifth Third (ex-Comerica)','G')]
        for bank,number in debts:self.assertEqual(section(self.model(bank,number)),'Debts')
        for bank,number in deposits:self.assertEqual(section(self.model(bank,number)),'Deposits')

    def test_separate_signed_totals_and_invalid_or_missing_values_excluded(self):
        rows=[self.model('CNB','D','100.25'),self.model('AMEX','A','20.50','LIABILITY CONVENTION UNVERIFIED'),self.model('Sapphire Chase','C','-1.25','LIABILITY CONVENTION UNVERIFIED')]
        invalid=self.model('AMEX','bad','999','AMBIGUOUS SETUP IDENTITY');rows.append(invalid)
        missing=self.model('CNB','missing');missing['Closing balance converted to USD']='NOT AVAILABLE';rows.append(missing)
        before=repr(rows)
        self.assertEqual(totals(rows),{'Deposits':Decimal('100.25'),'Debts':Decimal('19.25')})
        self.assertEqual(before,repr(rows))

    def test_both_html_views_and_excel_sections_use_identical_rows_dates_and_totals(self):
        rows=[self.model('CNB','D','100.25'),self.model('AMEX','A','20.50','LIABILITY CONVENTION UNVERIFIED')];before=repr(rows)
        now=datetime.fromisoformat('2026-10-22T12:00:00+03:00')
        full=print_document(rows,Decimal('100.25'),[],now);compact=compact_html(rows,now)
        for html in (full,compact):
            for name in ('Deposits','Debts'):self.assertIn(name,html)
            for amount in ('100.25','20.50'):self.assertIn(amount,html)
            self.assertIn('Balance Import Date',html.replace('<br>',' '));self.assertIn('Transaction Import Date',html.replace('<br>',' '))
        book=load_workbook(BytesIO(workbook_bytes(rows,Decimal('100.25'),[],now)))
        self.assertIn('Latest Import Balances',book.sheetnames)
        for name,group in partition(rows).items():
            sheet=book[name];self.assertEqual([c.value for c in sheet[1]],list(COLUMNS))
            self.assertEqual(sheet.cell(2,COLUMNS.index('Balance Import Date')+1).value,group[0]['Balance Import Date'])
            self.assertEqual(Decimal(sheet.cell(3,2).value),totals(rows)[name])
        self.assertEqual(before,repr(rows))

    def test_snapshot_selection_fields_and_fx_identical_to_live_only_bucket_total_changes(self):
        self.seed('deposit','2026-10-01T08:00:00+00:00',True)
        with closing(db.get_connection()) as c:
            c.execute("INSERT INTO account_list(account_name,bank,account_number,currency,rate_type) VALUES('Synthetic debt','Comerica','055-9000-0991 Credit line','USD','USD/USD')");c.commit()
        self.seed('debt','2026-10-02T08:00:00+00:00',True,number='055-9000-0991 Credit line',bank='Comerica')
        # Match source ownership, not a label-only correction in the report.
        with closing(db.get_connection()) as c:c.execute("UPDATE statement_balances SET account_name='Synthetic debt' WHERE statement_hash='debt'");c.execute("UPDATE classified_transactions SET account_name='Synthetic debt' WHERE statement_hash='debt'");c.commit()
        baseline=Path(__file__).resolve().parent.parent/'areti-urgent-import-family'/'latest_import_balances.py'
        spec=importlib.util.spec_from_file_location('pinned_live_balance_snapshot',baseline);prior=importlib.util.module_from_spec(spec);spec.loader.exec_module(prior)
        before=self.dump();old_rows,old_total,_=prior.snapshot(db);new_rows,new_total,_=snapshot(db)
        self.assertEqual(old_rows,new_rows)
        self.assertEqual(old_total,new_total+totals(new_rows)['Debts'])
        self.assertEqual(new_total,totals(new_rows)['Deposits']);self.assertEqual(before,self.dump())

if __name__=='__main__':
    names=[n for n in SectionTests.__dict__ if n.startswith('test_')]
    raise SystemExit(not unittest.TextTestRunner(verbosity=2).run(unittest.TestSuite(SectionTests(n) for n in names)).wasSuccessful())
