"""Source-bound BOC and processing-state checks; originals remain outside Git."""
from contextlib import closing
from decimal import Decimal
from io import BytesIO
import json, os, unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import pandas as pd
import parsing, db, streamlit as st
from streamlit.delta_generator import DeltaGenerator
from boc_import import BOCParseError, parse_pages, validate_preview
from import_history import commit_statement
from _qa_import_button_state import ButtonTests
from _qa_post_import_message import CompletionTests


def fixture(movements=(), opening='100.00', total=None):
    bban='00200195'+'0'*16;check=98-int(bban+'123400')%97
    text=f'BCYPCY2N\nAccount Number 000000000000\nIBAN CY{check:02d}{bban}\nCurrency EUR\nStatement Period: 01/02/2027 - 28/02/2027\nPage 1 / 1'
    words=[]
    def word(s,x,y):words.append(dict(text=str(s),x0=x,x1=x+25,top=y,upright=True))
    for s,x in [('Transaction',40),('Value',100),('TransactionDetails',145),('Debit',390),('Credit',460),('Balance',540)]:word(s,x,270)
    word('Balancebroughtforward',145,290);word(opening,540,290)
    running=Decimal(opening);debit=credit=Decimal(0)
    for i,(side,value,desc) in enumerate(movements):
        y=310+i*40;amount=Decimal(value)
        word('11/02/2027',40,y);word('11/02/2027',100,y);word(desc,145,y)
        word(value,390 if side=='debit' else 460,y)
        if side=='debit':debit+=amount;running-=amount
        else:credit+=amount;running+=amount
        word(f'{running:.2f}',540,y)
    y=310+len(movements)*40;word('Total / Balance Carried Forward',40,y)
    for x,value in zip((390,460,540),total or (debit,credit,running)):word(f'{Decimal(value):.2f}',x,y)
    return SimpleNamespace(extract_words=lambda:words),text


class ProcessingTests(ButtonTests):
    def test_processing_is_visible_before_preview_preparation(self):
        app=self.run_app(self.app());events=[];warning=DeltaGenerator.warning
        def warn(slot,body,*args,**kwargs):events.append(str(body));return warning(slot,body,*args,**kwargs)
        # A fresh file forces the actual preparation path rather than a cache hit.
        self.file=self.csv('second-preview.csv',[('2026-09-12','Synthetic preparation','-1.00')])
        app.session_state['statement_import_request']=__import__('db').build_statement_hash(self.file.getvalue())
        original=parsing.parse_csv
        def parse(*args,**kwargs):
            self.assertTrue(any('IN PROGRESS' in value for value in events));return original(*args,**kwargs)
        with patch.object(DeltaGenerator,'warning',warn),patch.object(parsing,'parse_csv',side_effect=parse):self.run_app(app)
        self.complete(app,1)
    def test_top_processing_replaces_preview_before_commit(self):
        app=self.run_app(self.app());events=[]
        warning=DeltaGenerator.warning;success=DeltaGenerator.success
        def warn(slot,body,*args,**kwargs):events.append(('warning',str(body)));return warning(slot,body,*args,**kwargs)
        def ok(slot,body,*args,**kwargs):events.append(('success',str(body)));return success(slot,body,*args,**kwargs)
        def commit(*args,**kwargs):
            self.assertTrue(any('IN PROGRESS' in value for kind,value in events))
            self.assertFalse(any('nothing has been imported' in value for kind,value in events))
            self.assertFalse(any('have been imported into the database' in value for kind,value in events))
            return commit_statement(*args,**kwargs)
        with patch.object(DeltaGenerator,'warning',warn),patch.object(DeltaGenerator,'success',ok),patch('import_history.commit_statement',side_effect=commit):self.import_file(app)
        self.complete(app,2)
        self.assertFalse(any('IN PROGRESS' in e.value for e in app.warning))

    def test_duplicate_and_failure_remove_processing(self):
        app=self.run_app(self.app())
        with patch('import_history.commit_statement',return_value=(0,True,0)):self.import_file(app)
        self.assertFalse(any('IN PROGRESS' in e.value for e in app.warning))
        app=self.run_app(self.app())
        with patch.object(db,'save_statement_balance',side_effect=ValueError('Synthetic failed commit')):self.import_file(app)
        self.assertFalse(any('IN PROGRESS' in e.value for e in app.warning))
        self.assertTrue(any('Synthetic failed commit' in e.value for e in app.error))
        self.assertFalse(next(b for b in app.button if b.label=='Import statement').disabled)


class ColumnTests(unittest.TestCase):
    def test_labels_override_misleading_words(self):
        page,text=fixture([('credit','17.43','CARD FEES PURCHASE'),('debit','7.43','INWARD CREDIT REFUND')])
        rows,meta=parse_pages([page],[text]);self.assertEqual([r[2] for r in rows],[Decimal('17.43'),Decimal('-7.43')])
        self.assertEqual(meta['closing_balance'],Decimal('110.00'))

    def test_zero_row_explicit_totals(self):
        page,text=fixture();rows,meta=parse_pages([page],[text]);self.assertEqual(rows,[])
        self.assertTrue(meta['zero_activity_validated']);self.assertEqual(meta['money_in'],0);self.assertEqual(meta['money_out'],0)

    def test_fail_closed_layout_and_arithmetic(self):
        for mode in ('total','running','header','direction','negative','truncated','iban','unknown'):
            page,text=fixture([('credit','17.43','Synthetic')]);words=page.extract_words()
            if mode=='total':next(w for w in words if w['top']==350 and w['x0']==460)['text']='18.43'
            if mode=='running':next(w for w in words if w['top']==310 and w['x0']==540)['text']='82.57'
            if mode=='header':words[:]=[w for w in words if w['text']!='Credit']
            if mode=='direction':next(w for w in words if w['top']==310 and w['x0']==460)['x0']=390
            if mode=='negative':next(w for w in words if w['top']==310 and w['x0']==460)['text']='-17.43'
            if mode=='truncated':text=text.replace('1 / 1','1 / 2')
            if mode=='iban':text=text.replace('IBAN CY','IBAN XX')
            if mode=='unknown':words.append(dict(text='UNKNOWN',x0=145,x1=200,top=300))
            with self.subTest(mode=mode),self.assertRaises(BOCParseError):parse_pages([page],[text])

    def test_preview_mutation_rejected_before_database_access(self):
        page,text=fixture([('credit','17.43','Synthetic')]);rows,meta=parse_pages([page],[text])
        account=dict(account_number=meta['account_number'],currency='EUR')
        frame=pd.DataFrame({'Date':[rows[0][0]],'Amount':[rows[0][2]],'account_number':[account['account_number']],'currency':['EUR']})
        validate_preview(frame,meta,account)
        for mode in ('amount','date','currency','account','missing','balance'):
            changed=frame.copy();balance=dict(meta);selected=dict(account)
            if mode=='amount':changed.loc[0,'Amount']=Decimal('-17.43')
            if mode=='date':changed.loc[0,'Date']='2027-02-12'
            if mode=='currency':changed.loc[0,'currency']='USD'
            if mode=='account':selected['account_number']='9999'
            if mode=='missing':changed=changed.iloc[:0]
            if mode=='balance':balance['closing_balance']=Decimal('118.43')
            with self.subTest(mode=mode),patch.object(db,'get_connection',side_effect=AssertionError('DB opened before validation')),patch.object(db,'statement_already_imported',return_value=False),self.assertRaises(BOCParseError):
                commit_statement(db,changed,'synthetic.pdf','synthetic',balance,selected)


class SourceTests(CompletionTests):
    def setUp(self):
        super().setUp()
        self.sources=json.loads((Path(os.environ['BOC_PACKAGE_PRIVATE'])/'baseline-probe.json').read_text())
        # Disposable synthetic Setup rate; never read a production exchange rate.
        with closing(db.get_connection()) as conn:
            conn.execute("INSERT INTO rates(rate_month,rate_type,rate_value) VALUES('2026-09-01','EUR/USD',1.17)")
            conn.commit()

    def test_all_original_rows_source_columns_and_metadata(self):
        expected=json.loads((Path(os.environ['BOC_PACKAGE_PRIVATE'])/'source-expected.json').read_text())
        for source,amounts,balances,dates in zip(self.sources,expected['expected_amounts'],expected['expected_balances'],expected['expected_dates']):
            path=Path(source['source']);frame=parsing.parse_pdf(BytesIO(path.read_bytes()));meta=frame.attrs['statement_balance']
            self.assertEqual(frame.Amount.tolist(),list(map(Decimal,amounts)))
            self.assertEqual(frame.Date.tolist(),dates)
            self.assertEqual(tuple(meta[k] for k in ('opening_balance','money_in','money_out','closing_balance')),tuple(map(Decimal,balances)))
            self.assertEqual(meta['period_start'],'2026-09-01');self.assertEqual(meta['period_end'],'2026-09-30')
            self.assertEqual(meta['account_number'],source['meta']['account_number']);self.assertEqual(meta['iban'],source['meta']['iban'])
            self.assertEqual(meta['opening_balance']+meta['money_in']-meta['money_out'],meta['closing_balance'])
            self.assertEqual(parsing.extract_statement_balance(BytesIO(path.read_bytes()),path.name),meta)

    def test_upload_atomic_history_pending_latest_balances_duplicates(self):
        from latest_import_balances import snapshot,print_document,workbook_bytes
        from openpyxl import load_workbook
        for source in self.sources:
            path=Path(source['source']);meta=source['meta']
            with closing(db.get_connection()) as conn:
                conn.execute('INSERT INTO account_list(account_name,bank,account_number,currency,rate_type) VALUES(?,?,?,?,?)',('Source QA','Bank of Cyprus',meta['account_number'],meta['currency'],'USD/USD' if meta['currency']=='USD' else 'EUR/USD'))
                conn.commit()
            st.cache_data.clear();self.file=BytesIO(path.read_bytes());self.file.name=path.name
            before=self.snapshot();app=self.run_app(self.app());self.assertEqual(self.snapshot(),before)
            with patch.object(db,'save_statement_balance',side_effect=ValueError('BOC atomic rollback')):self.import_file(app)
            self.assertEqual(self.snapshot(),before)
            self.import_file(app);self.complete(app,meta['transaction_count'])
            pending=db.get_pending_transactions();matched=pending[pending.account_number==meta['account_number']]
            self.assertEqual(len(matched),meta['transaction_count'])
            self.assertEqual(sorted(Decimal(str(value)) for value in matched.amount),sorted(Decimal(row[2]) for row in source['rows']))
            history=db.get_import_history();h=history[history.account_number==meta['account_number']].iloc[0]
            self.assertEqual(h.reconciliation_status,'OK');self.assertEqual(Decimal(str(h.closing_balance)),Decimal(meta['closing_balance']))
            report,total,warnings=snapshot(db);row=next(row for row in report if row['Account number']==meta['account_number'])
            self.assertEqual(Decimal(str(row['Closing balance'])),Decimal(meta['closing_balance']))
            self.assertEqual(Decimal(str(row['Opening balance'])),Decimal(meta['opening_balance']))
            self.assertEqual(row['Statement start date'],'2026-09-01');self.assertEqual(row['Statement end date'],'2026-09-30')
            self.assertIn('INTERIM PARTIAL REPORT',print_document(report,total,warnings))
            workbook=load_workbook(BytesIO(workbook_bytes(report,total,warnings)),data_only=True)
            exported=next(values for values in workbook['Latest Import Balances'].iter_rows(min_row=2,values_only=True) if values[0]==meta['account_number'])
            self.assertEqual(Decimal(str(exported[7])),Decimal(meta['closing_balance']))
            self.assertIn('INTERIM PARTIAL REPORT',workbook['Verification']['B2'].value)
            before=self.snapshot();self.run_app(app);self.assertEqual(self.snapshot(),before)
            self.assertFalse(any(b.label=='Import statement' for b in app.button))
            other=self.run_app(self.app());self.assertTrue(any('already exists' in e.value for e in other.warning));self.assertEqual(self.snapshot(),before)
        # Actual page dispatch after atomic imports, without uploading again.
        app.query_params['page']='Latest Import Balances'
        app.session_state['main_navigation']='Latest Import Balances'
        with patch.object(st,'file_uploader',return_value=None):app.run()
        self.assertFalse(app.exception);self.assertFalse(app.error)
        self.assertIn('Latest Import Balances',[e.value for e in app.subheader])
        self.assertEqual(self.snapshot(),before)
        self.assertTrue(any('INTERIM PARTIAL REPORT' in e.value for e in app.info))

    def test_zero_row_boc_completion_and_duplicate(self):
        page,text=fixture();rows,meta=parse_pages([page],[text])
        frame=parsing._frame_from_pdf_rows(rows);frame.attrs['statement_balance']=meta
        with closing(db.get_connection()) as conn:
            conn.execute('INSERT INTO account_list(account_name,bank,account_number,currency,rate_type) VALUES(?,?,?,?,?)',('Zero source QA','Bank of Cyprus',meta['account_number'],'EUR','EUR/USD'));conn.commit()
        st.cache_data.clear();self.file=BytesIO(b'validated synthetic zero-row fixture');self.file.name='zero-boc.csv'
        with patch.object(parsing,'parse_csv',return_value=frame):
            app=self.run_app(self.app());before=self.snapshot();self.import_file(app)
            self.assertFalse(app.error,[e.value for e in app.error]);self.assertEqual(self.counts(),(0,1,1))
            self.assertTrue(any('No new transactions were added' in e.value for e in app.success))
            self.assertFalse(any('nothing has been imported' in e.value for e in app.success))
            after=self.snapshot();self.run_app(app);self.assertEqual(self.snapshot(),after)
            self.assertFalse(any(b.label=='Import statement' for b in app.button))


if __name__=='__main__':unittest.main(verbosity=2)
