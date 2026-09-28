"""Synthetic regressions for monetary importer and presentation boundaries."""
from decimal import Decimal
from io import BytesIO


def main():
    import parsing
    import utils
    import reporting
    from _qa_revolut_business import TEXT
    value = Decimal('90071992547409.37')
    assert utils.format_currency(value) == '90,071,992,547,409.37', 'Formatter changed cents'
    assert reporting._format_money(value) == '90,071,992,547,409.37', 'PDF formatter changed cents'
    rows = parsing._parse_revolut_business_pdf_text(TEXT)
    assert rows and all(isinstance(row[2], Decimal) for row in rows), 'Importer converted exact amounts to float'
    assert utils.format_currency(Decimal('-0.01')) == '-0.01'
    assert utils.format_currency(Decimal('0.00')) == '0.00'
    assert utils.format_currency(None) == '0.00'
    import openpyxl
    from financial_decimal import excel_value
    book = openpyxl.Workbook()
    book.active.cell(1, 1, excel_value(value))
    stream = BytesIO()
    book.save(stream)
    book.close()
    loaded = openpyxl.load_workbook(BytesIO(stream.getvalue()), read_only=True)
    assert Decimal(loaded.active.cell(1, 1).value) == value
    loaded.close()
    from financial_tabular import read_excel_exact
    from zipfile import ZipFile
    from xml.etree import ElementTree
    source = openpyxl.Workbook()
    source.active.append(['Date', 'Description', 'Amount'])
    source.active.append(['2026-01-10', 'Synthetic', 1])
    raw = BytesIO()
    source.save(raw)
    source.close()
    exact_source = BytesIO()
    with ZipFile(BytesIO(raw.getvalue())) as original, ZipFile(exact_source, 'w') as modified:
        for member in original.infolist():
            data = original.read(member.filename)
            if member.filename == 'xl/worksheets/sheet1.xml':
                root = ElementTree.fromstring(data)
                ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                cell = root.find('.//s:c[@r="C2"]/s:v', ns)
                cell.text = str(value)
                data = ElementTree.tostring(root)
            modified.writestr(member, data)
    assert read_excel_exact(exact_source).Amount.iloc[0] == value
    assert parsing.parse_excel(exact_source).Amount.iloc[0] == value
    csv = BytesIO(('Date,Description,Amount\n2026-01-10,Synthetic,' + str(value) + '\n').encode())
    assert parsing.parse_csv(csv).Amount.iloc[0] == value
    for text in ('1.234567890123456789E-8', '-9.007199254740937E13', '0E-8'):
        scientific = BytesIO(('Date,Description,Amount\n2026-01-10,Synthetic,' + text + '\n').encode())
        assert parsing.parse_csv(scientific).Amount.iloc[0] == Decimal(text), 'Scientific CSV source value changed'
    import analytics
    import pandas as pd
    import db
    serial_date = Decimal((pd.Timestamp('2026-01-01') - pd.Timestamp('1899-12-30')).days)
    assert db._parse_rate_month_cell(serial_date) == pd.Timestamp('2026-01-01'), 'Decimal Excel date ordinal was lost'
    assert pd.isna(db._parse_rate_month_cell(Decimal('1.234567890123456789'))), 'An FX rate became a date'
    frame = pd.DataFrame([dict(id=i, txn_date=date, amount=-value, usd_amount=-value,
        category='Synthetic', normalized_description='SYNTHETIC', original_description='Synthetic')
        for i, date in enumerate(['2026-01-10', '2026-02-10'], 1)])
    context = analytics.build_report_context(frame, 2, 'All')
    assert context['kpis']['total_expenses'] == value * 2
    assert isinstance(context['kpis']['total_expenses'], Decimal)
    assert context['prediction'] == value
    analytics.detect_anomalies(frame)
    import ast
    from pathlib import Path
    tree = ast.parse(Path('app.py').read_text(encoding='utf-8'))
    assignment_nodes = [n for n in ast.walk(tree) if isinstance(n, ast.Assign)
        and any(isinstance(t, ast.Subscript) and isinstance(t.value, ast.Name)
                and t.value.id == 'filtered_reviewed' and isinstance(t.slice, ast.Constant)
                and t.slice.value in ('amount', 'amount_usd') for t in n.targets)]
    from financial_decimal import optional_decimal
    environment = dict(pd=pd, optional_decimal=optional_decimal,
        filtered_reviewed=pd.DataFrame({'amount': [value], 'amount_usd': [value]}))
    assert len(assignment_nodes) == 2
    exec(compile(ast.Module(body=assignment_nodes, type_ignores=[]), 'app.py', 'exec'), environment)
    assert environment['filtered_reviewed'].amount.iloc[0] == value
    print('PASS exact importer amounts, large/negative/zero/NULL formatting')


if __name__ == '__main__':
    main()
