"""Actual local PostgreSQL import -> reporting -> exact-text export contract."""
from contextlib import closing
from decimal import Decimal
from io import BytesIO
import os
from pathlib import Path
from unittest.mock import patch
import uuid


def main():
    for key in ('DATABASE_URL', 'POSTGRES_URL'):
        os.environ.pop(key, None)
    import openpyxl
    import pandas as pd
    import psycopg2
    import db
    import reporting
    from financial_schema import FINANCIAL_COLUMNS
    from financial_decimal import product, cents

    options = dict(host='127.0.0.1', port=int(os.environ['ARETI_QA_PG_PORT']),
                   user='qa_local', sslmode='disable')
    name = 'qa_financial_pipeline_' + uuid.uuid4().hex
    with closing(psycopg2.connect(dbname='postgres', **options)) as admin:
        admin.autocommit = True
        with admin.cursor() as cur:
            cur.execute('SHOW data_directory')
            assert Path(cur.fetchone()[0]).resolve() == Path(os.environ['ARETI_QA_PG_DATA']).resolve()
            cur.execute('CREATE DATABASE ' + name)
        def connect():
            return db.PostgresConnection(psycopg2.connect(dbname=name, **options))
        try:
            with patch.object(db, 'USING_POSTGRES', True), patch.object(db, 'get_connection', connect):
                db.init_db()
                with closing(connect()) as connection:
                    for table, columns in FINANCIAL_COLUMNS.items():
                        for field in columns:
                            connection.cursor().execute(f'ALTER TABLE {table} ALTER COLUMN {field} TYPE NUMERIC')
                    connection.commit()
                db.add_category('Synthetic', 'General', 'Family expenses')
                values = [Decimal('-90071992547409.37'), Decimal('-712345.30')]
                rate = Decimal('1.234567890123456789')
                rate_book = openpyxl.Workbook()
                rate_book.active.append(['Rate', '2026-01-01'])
                rate_book.active.append(['GBP/USD', str(rate)])
                rate_upload = BytesIO()
                rate_book.save(rate_upload)
                rate_book.close()
                assert db.replace_rates_from_excel(rate_upload) == 1
                assert db.get_rates().iloc[0].rate_value == rate, 'Rate spreadsheet lost precision'
                balance = dict(bank='Synthetic', account_number='SYNTHETIC', currency='GBP',
                    period_start='2026-01-01', period_end='2026-01-31',
                    opening_balance=Decimal('90071992547409.37'), money_in=Decimal('0.00'),
                    money_out=Decimal('0.00'), closing_balance=Decimal('90071992547409.37'))
                db.save_statement_balance('synthetic-exact', 'synthetic.pdf', balance)
                balances = db.get_statement_balances()
                assert balances.iloc[0].opening_balance == balance['opening_balance'], 'Balance read lost precision'
                frame = pd.DataFrame([dict(Date='2026-01-15', Description='Synthetic exact ' + str(i),
                    Amount=amount, currency='GBP', amount_usd=cents(product(amount, rate)), fx_rate=rate,
                    suggested_category='Synthetic', suggested_subcategory='General')
                    for i, amount in enumerate(values)])
                inserted, duplicate, _ = db.save_pending_transactions(frame, 'synthetic.pdf', 'synthetic-exact')
                assert inserted == 2 and not duplicate
                history = db.get_import_history()
                assert len(history) == 1
                assert history.iloc[0].opening_balance == balance['opening_balance'], 'History read lost precision'
                saved = db.get_all_transactions().sort_values('id')
                assert list(saved.amount) == values
                assert [str(v) for v in saved.amount] == [str(v) for v in values], 'Stored scale changed'
                assert all(isinstance(v, Decimal) for v in saved.amount)
                assert list(saved.fx_rate) == [rate, rate]
                _, _, summary, _, _ = reporting._build_sections(saved, db.get_categories(include_subcategories=True))
                expected = sum((cents(product(v, rate)).copy_abs() for v in values), Decimal(0))
                assert sum((Decimal(str(r['total'])) for r in summary), Decimal(0)) == expected, 'Report changed exact imported cents'
                exported = db.dataframe_to_excel_bytes({'Amounts': saved[['amount', 'fx_rate']]})
                workbook = openpyxl.load_workbook(BytesIO(exported), read_only=True)
                try:
                    rows = list(workbook.active.values)[1:]
                    assert [Decimal(r[0]) for r in rows] == values
                    assert all(Decimal(r[1]) == rate for r in rows)
                finally:
                    workbook.close()
                report = reporting.build_sample_expenses_report(saved, db.get_categories(include_subcategories=True))
                workbook = openpyxl.load_workbook(BytesIO(report), read_only=True)
                try:
                    total_rows = [r for r in workbook['Sample expenses report'].values if r[0] == 'TOTAL']
                    assert len(total_rows) == 1 and Decimal(str(total_rows[0][1])) == expected, 'Report XLSX total lost precision'
                    exported_rows = list(workbook['Reviewed transactions'].values)
                    amount_column = exported_rows[0].index('amount')
                    assert sorted(Decimal(str(r[amount_column])) for r in exported_rows[1:]) == sorted(values)
                finally:
                    workbook.close()
                print('PASS actual import/database/report/export exact roundtrip, large negative amounts and fine FX')
        finally:
            with admin.cursor() as cur:
                cur.execute('DROP DATABASE ' + name)


if __name__ == '__main__':
    main()
