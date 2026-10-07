"""Focused metadata display/export checks; disposable databases only."""
from contextlib import closing
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
import unittest
import db
from _qa_balance_cnb_report import BalanceTests
from import_history import commit_statement
from latest_import_balances import snapshot, print_document, workbook_bytes, COLUMNS, NOTE
from openpyxl import load_workbook


class FlowDisplayTests(BalanceTests):
    def test_stored_flows_visible_in_report_and_exports_without_mutation(self):
        frame, balance, account = self.source()
        commit_statement(db, frame, 'synthetic-flow.pdf', 'flow', balance, account)
        before = self.snapshot()
        rows, total, warnings = snapshot(db)
        row = next(row for row in rows if row['Account number'] == account['account_number'])
        self.assertEqual(row['Credits / money in'], Decimal('10.00'))
        self.assertEqual(row['Debits / money out'], Decimal('0.00'))
        self.assertEqual(row['Closing balance'], Decimal('110.00'))
        html = print_document(rows, total, warnings)
        self.assertIn('Credits / money in', html)
        self.assertIn('Debits / money out', html)
        self.assertIn(NOTE, html)
        book = load_workbook(BytesIO(workbook_bytes(rows, total, warnings)))
        sheet = book['Latest Import Balances']
        headings = [cell.value for cell in sheet[1]]
        actual = next(values for values in sheet.iter_rows(min_row=2, values_only=True)
                      if values[0] == account['account_number'])
        self.assertEqual(actual[headings.index('Credits / money in')], 10)
        self.assertEqual(actual[headings.index('Debits / money out')], 0)
        self.assertEqual(len(headings), len(COLUMNS))
        self.assertEqual(self.snapshot(), before)

    def test_unknown_flows_remain_unknown(self):
        frame, balance, account = self.source()
        commit_statement(db, frame, 'synthetic-legacy.pdf', 'legacy', balance, account)
        with closing(db.get_connection()) as conn:
            conn.execute("UPDATE statement_balances SET money_in=NULL,money_out=NULL WHERE statement_hash='legacy'")
            conn.commit()
        before = self.snapshot()
        rows, total, warnings = snapshot(db)
        row = next(row for row in rows if row['Account number'] == account['account_number'])
        self.assertIsNone(row['Credits / money in'])
        self.assertIsNone(row['Debits / money out'])
        print_document(rows, total, warnings)
        sheet = load_workbook(BytesIO(workbook_bytes(rows, total, warnings)))['Latest Import Balances']
        headings = [cell.value for cell in sheet[1]]
        actual = next(values for values in sheet.iter_rows(min_row=2, values_only=True)
                      if values[0] == account['account_number'])
        self.assertIsNone(actual[headings.index('Credits / money in')])
        self.assertIsNone(actual[headings.index('Debits / money out')])
        self.assertEqual(self.snapshot(), before)


if __name__ == '__main__':
    unittest.main()
