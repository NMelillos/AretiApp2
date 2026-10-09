import os,sys,unittest
from pathlib import Path
sys.path.insert(0,str(Path.cwd()))
from datetime import datetime,timedelta
from zoneinfo import ZoneInfo
from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook
from latest_import_balances import fresh_closing,print_document,workbook_bytes,COLUMNS
from latest_balances_compact import compact_html,compact_model
from _qa_post_import_message import CompletionTests
import db,pandas as pd
from contextlib import closing

class ReportTests(unittest.TestCase):
 def test_closing_age_boundaries(self):
  now=datetime(2026,10,9,0,1,tzinfo=ZoneInfo('Europe/Nicosia'))
  for age in (-1,0,29,30,31):self.assertEqual(fresh_closing((now.date()-timedelta(days=age)).isoformat(),now),0<=age<=30)
  for value in ('',None,'bad','2026-02-30'):self.assertFalse(fresh_closing(value,now))
 def test_exact_visible_total_columns_sort_and_no_grouping(self):
  rows=[]
  for bank,stamp,end,value in [('Z Bank','2026-10-09 10:00:00 EEST','2026-08-01','12.340000'),('A Bank','2026-10-08 10:00:00 EEST','2026-09-30','2.10'),('A Bank','2026-10-07 10:00:00 EEST','2026-09-30','3.20')]:
   row={k:'' for k in COLUMNS};row.update({'Bank':bank,'Import date':stamp,'Statement end date':end,'Account number':'0001','Account name':'<source>','Currency':'USD','Closing balance':Decimal(value),'Closing balance converted to USD':Decimal(value),'Status':'IMPORTED','Verification':'SOURCE RECONCILIATION NOT VERIFIED'});rows.append(row)
  visible,total,*_=compact_model(rows);self.assertEqual(total,Decimal('17.64'));self.assertEqual([r['import_date'] for r in visible[:2]],[rows[1]['Import date'],rows[2]['Import date']])
  now=datetime(2026,10,9,tzinfo=ZoneInfo('Europe/Nicosia'));html=compact_html(rows,now)
  self.assertIn('colspan="7"',html);self.assertIn('17.64',html);self.assertNotIn('12.340000',html);self.assertIn('&lt;source&gt;',html);self.assertEqual(html.count('<tr>'),5);self.assertIn('color:#b42318',html);self.assertIn('color:#146b36',html);self.assertIn('color:#000000',html)
  doc=print_document(rows,total,[],now);self.assertNotIn('Under 30 days',doc)
  wb=load_workbook(BytesIO(workbook_bytes(rows,total,[],now)));sheet=wb.active
  self.assertEqual(sheet.cell(2,4).font.color.rgb,'00000000');self.assertEqual(sheet.cell(2,6).font.color.rgb,'00B42318');self.assertEqual(sheet.cell(3,6).font.color.rgb,'00146B36')
  rows[0]['Verification']='unverified';self.assertEqual(compact_model(rows)[1],Decimal('5.30'));self.assertIn('—',compact_html(rows))

class AccountDeletionTest(CompletionTests):
 def test_setup_removal_preserves_historical_records(self):
  app=self.run_app(self.app());self.import_file(app)
  tables=('classified_transactions','transaction_memory','statement_imports','statement_balances')
  def snapshot():
   with closing(db.get_connection()) as conn:return [conn.execute('SELECT * FROM '+t+' ORDER BY id').fetchall() for t in tables]
  # Retain actual reviewed classifications and non-empty learned memory too.
  with closing(db.get_connection()) as conn:
   conn.execute("UPDATE classified_transactions SET category='Synthetic',subcategory='General',reviewed=1 WHERE id=(SELECT min(id) FROM classified_transactions)")
   conn.execute("INSERT INTO transaction_memory(normalized_description,category,subcategory) VALUES('synthetic remembered','Synthetic','General')")
   conn.commit()
  before=snapshot();output=BytesIO()
  pd.DataFrame([dict(**{'Account Name':'Retained','Bank':'Test','Account Number':'0002','Currency':'USD','Rate Type':'USD/USD'})]).to_excel(output,index=False)
  output.seek(0);db.replace_accounts_from_excel(output)
  self.assertEqual(snapshot(),before);self.assertNotIn('Synthetic',db.get_accounts().account_name.tolist());self.assertEqual(len(db.get_pending_transactions()),1)

if __name__=='__main__':
 suite=unittest.TestSuite([unittest.defaultTestLoader.loadTestsFromTestCase(ReportTests),AccountDeletionTest('test_setup_removal_preserves_historical_records')])
 result=unittest.TextTestRunner(verbosity=2).run(suite);sys.exit(not result.wasSuccessful())
