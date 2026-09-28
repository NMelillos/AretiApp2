"""Lossless, authenticated, in-memory XLSX export of the displayed snapshot."""
from contextlib import ExitStack
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import re
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook


def verify_workbook(data, diagnostics, comparison):
    book = load_workbook(BytesIO(data), data_only=False)
    expected = {
        'Record diagnostics': diagnostics['fields'],
        'Schema metadata': diagnostics['schema'],
        'Precondition hashes': diagnostics['hashes'],
        'Account sections': comparison['sections'],
        'Transactions': comparison['transactions'],
    }
    assert book.sheetnames == list(expected) + ['Export metadata']
    for name, rows in expected.items():
        sheet = book[name]
        columns = list(dict.fromkeys(k for row in rows for k in row))
        actual = list(sheet.values)
        wanted = [tuple(columns)] + [tuple('NULL' if row.get(k) is None else str(row[k]) for k in columns) for row in rows]
        if rows:
            assert [tuple('' if v is None else v for v in row) for row in actual] == wanted, name
            assert sheet.freeze_panes == 'A2' and sheet.auto_filter.ref == sheet.dimensions
        for row in sheet:
            assert all(cell.data_type != 'f' for cell in row), 'Formula injection'
    metadata = dict(list(book['Export metadata'].values)[1:])
    assert metadata['PDF fingerprint'] == comparison['fingerprint']
    assert metadata['Related audit records'] == str(diagnostics['audit_count'])
    assert metadata['Mode'] == 'READ ONLY - NO REPAIR'
    with ZipFile(BytesIO(data)) as archive:
        assert not any(name.startswith('xl/externalLinks') for name in archive.namelist())
    book.close()


def main():
    import db
    import existing_import_compare as compare
    import repair_readiness as readiness
    comparison = {'fingerprint': 'synthetic-fingerprint',
                  'sections': [{'Import ID': 1, 'Balance ID': 2, 'PDF opening_balance': '10.00'}],
                  'transactions': [{'Import ID': 1, 'Record ID': 3, 'Stored amount': '-12.345678901234567890123456789'}]}
    diagnostics = {'fields': [], 'schema': [{'Table': 'classified_transactions', 'column_name': 'amount',
        'data_type': 'numeric', 'numeric_precision': None, 'numeric_scale': None, 'numeric_precision_radix': 10}],
        'hashes': [{'Scope': 'classified_transactions', 'Record ID': '3', 'Precondition SHA-256': 'a' * 64}],
        'audit_count': 0, 'not_displayed': 'must-not-export'}
    values = ['-12.345678901234567890123456789', '12.30000000000000000000', '00001234', '',
              '=HYPERLINK("https://example.invalid","synthetic")', '+1+1', '-1+1', '@SUM(A1)',
              'Synthetic\ttext\nsecond line', '\u0394\u03bf\u03ba\u03b9\u03bc\u03ae', '#N/A', 'NULL']
    for index in range(360):
        diagnostics['fields'].append({'Table': 'classified_transactions', 'Record ID': index + 1,
            'Field': 'amount', 'Current database value': values[index % len(values)],
            'Proposed PDF/parser value': Decimal('12.30000000000000000000'),
            'Disposition': 'UNCHANGED / PRESERVE'})
    original = deepcopy((diagnostics, comparison))
    with patch.object(compare, 'get_script_run_ctx', return_value=object()), patch.object(
            compare.st, 'session_state', {'authenticated': True, 'login_user': 'Areti'}), patch.object(
            db, 'get_connection', side_effect=AssertionError('Export must not query the database')):
        data, filename = readiness.export_workbook(diagnostics, comparison)
        assert re.fullmatch(r'NOMAD_Repair_Diagnostics_\d{8}_\d{6}_UTC\.xlsx', filename)
        verify_workbook(data, diagnostics, comparison)
        assert (diagnostics, comparison) == original
        assert b'must-not-export' not in ZipFile(BytesIO(data)).read('xl/worksheets/sheet6.xml')
        for bad in ('x' * 32768, 'invalid\x00control'):
            altered = deepcopy(diagnostics)
            altered['fields'][0]['Current database value'] = bad
            try:
                readiness.export_workbook(altered, comparison)
            except ValueError:
                pass
            else:
                raise AssertionError('Invalid/overlong value silently changed')
    for state in ({}, {'authenticated': True, 'login_user': 'areti'},
                  {'authenticated': True, 'login_user': 'Other'},
                  {'authenticated': 'true', 'login_user': 'Areti'},
                  {'authenticated': True, 'login_user': 'Areti', 'third_report_authenticated': True}):
        with patch.object(compare, 'get_script_run_ctx', return_value=object()), patch.object(
                compare.st, 'session_state', state):
            try:
                readiness.export_workbook(diagnostics, comparison)
            except compare.CompareBlocked:
                pass
            else:
                raise AssertionError('Unauthorized export')
    with patch.object(compare, 'get_script_run_ctx', return_value=None), patch.object(
            compare.st, 'session_state', {'authenticated': True, 'login_user': 'Areti'}):
        try:
            readiness.export_workbook(diagnostics, comparison)
        except compare.CompareBlocked:
            pass
        else:
            raise AssertionError('Export outside main application session')
    print('PASS lossless XLSX: all rows/columns, precision, identifiers, schema, links, hashes, formula safety, no DB access, authorization')


if __name__ == '__main__':
    main()
