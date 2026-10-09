"""Synthetic storage provenance, liability, legacy time and export regressions."""
from contextlib import closing
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
import os
from pathlib import Path
import tempfile
from unittest.mock import patch
import uuid

from openpyxl import load_workbook
import db
from latest_import_balances import snapshot, print_document, workbook_bytes, fresh_import, fresh_closing


def main():
    assert not db.USING_POSTGRES
    with tempfile.TemporaryDirectory(dir=os.environ['TEMP']) as root, patch.object(db,'DB_PATH',str(Path(root)/'report.sqlite')):
        db.init_db()
        with closing(db.get_connection()) as conn:
            schema=conn.execute("SELECT sql FROM sqlite_master WHERE name='statement_balances'").fetchone()[0]
            conn.execute('DROP TABLE statement_balances')
            conn.execute(schema.replace('opening_balance REAL','opening_balance TEXT').replace('closing_balance REAL','closing_balance TEXT'))
            def statement(index,amount,bank='Synthetic deposit',timestamp='2026-10-04T08:00:00+00:00'):
                key=f'synthetic-{index}'
                conn.execute('INSERT INTO account_list(account_name,bank,account_number,currency,rate_type) VALUES(?,?,?,?,?)',
                             (f'Synthetic {index}',bank,f'QA-{index}','USD','USD/USD'))
                conn.execute('INSERT INTO statement_imports(statement_hash,imported_at,transaction_count) VALUES(?,?,0)',(key,timestamp))
                conn.execute('INSERT INTO statement_balances(statement_hash,account_name,bank,account_number,currency,period_start,period_end,opening_balance,closing_balance) VALUES(?,?,?,?,?,?,?,?,?)',
                             (key,f'Synthetic {index}',bank,f'QA-{index}','USD','2026-09-01','2026-09-30',amount,amount))
            for index,amount in enumerate(('17.43','-7.19','0','17.43000030517578'),1):statement(index,amount)
            statement(5,'25.00','Citi')
            statement(6,None)
            statement(7,'1.01',timestamp='not-a-calendar-date')
            statement(8,'2.02',timestamp='2026-10-04 11:00:00')
            # Same actual identity/currency: latest timestamp then latest ID.
            conn.execute("INSERT INTO statement_imports(statement_hash,imported_at,transaction_count) VALUES('newest-seven','2026-10-05T08:00:00+00:00',0)")
            conn.execute("INSERT INTO statement_balances(statement_hash,account_name,bank,account_number,currency,period_start,period_end,opening_balance,closing_balance) VALUES('newest-seven','Synthetic 7','Synthetic deposit','QA-7','USD','2026-10-01','2026-10-31','3.03','3.03')")
            conn.execute("INSERT INTO account_list(account_name,bank,account_number,currency) VALUES('No import','Synthetic deposit','QA-NONE','USD')")
            conn.commit();before='\n'.join(conn.iterdump())
        rows,total,warnings=snapshot(db);by_account={row['Account number']:row for row in rows}
        assert total==Decimal('15.29')
        assert by_account['QA-4']['Verification']=='UNVERIFIED STORED PRECISION'
        assert by_account['QA-5']['Verification']=='LIABILITY CONVENTION UNVERIFIED'
        assert by_account['QA-5']['Closing balance converted to USD']==Decimal('25.00')
        assert by_account['QA-6']['Status']=='INCOMPLETE'
        assert by_account['QA-7']['Closing balance']==Decimal('3.03')
        assert any('ordering is unverified' in value for value in warnings)
        assert 'timezone unverified' in by_account['QA-8']['Import date']
        assert by_account['QA-NONE']['Status']=='NO IMPORT'
        generated=datetime.fromisoformat('2026-10-05T12:00:00+03:00')
        for days in (0,29,30,31,40,-1):
            stamp=(generated-timedelta(days=days)).strftime('%Y-%m-%d %H:%M:%S EEST')
            assert not fresh_import(stamp,generated)
            assert fresh_closing((generated-timedelta(days=days)).date().isoformat(),generated)==(0<=days<=30)
        for stamp in ('','invalid','2026-10-05 11:00:00 (timezone unverified)'):
            assert not fresh_import(stamp,generated)
        for date_value in ('2027-01-01T01:00:00+02:00','2026-03-30T01:00:00+03:00','2026-10-26T01:00:00+02:00'):
            current=datetime.fromisoformat(date_value)
            assert fresh_closing((current-timedelta(days=30)).date().isoformat(),current)
            assert not fresh_import((current-timedelta(days=29)).strftime('%Y-%m-%d %H:%M:%S EET'),current)
        rows[0]['Account name']='=HYPERLINK("https://example.invalid","unsafe")'
        rows[0]['Bank']='<script>unsafe</script>'
        html=print_document(rows,total,warnings,generated)
        assert '<script>unsafe</script>' not in html and 'PARTIAL TOTAL' in html
        assert 'current Setup accounts only' in html and 'not a complete reconciled total' in html
        assert 'TOTAL USD VALUE OF ALL ACCOUNTS' not in html
        assert 'Excluded balances and other verification warnings' in html
        assert '17.43' in html and '15.29' in html and 'Card/liability balances' in html
        book=load_workbook(BytesIO(workbook_bytes(rows,total,warnings,generated)))
        assert book.active['B2'].data_type=='s'
        assert book.active['G2'].number_format.startswith('#,##0.00')
        assert book['Verification']['B4'].value==float(total)
        assert 'current Setup accounts only' in book['Verification']['B2'].value
        assert 'Partial USD total' in book['Verification']['A4'].value
        assert any('QA-4' in str(row[1].value) for row in book['Verification'] if len(row)>1)
        assert any(row[0].value=='Verification warning' and 'ordering is unverified' in str(row[1].value) for row in book['Verification'])
        assert not any(row[0].value=='Excluded' and 'ordering is unverified' in str(row[1].value) for row in book['Verification'])
        complete_html=print_document([rows[0]],Decimal('17.43'),[],generated)
        assert 'PARTIAL TOTAL' in complete_html and 'outside Setup remain outside scope' in complete_html
        assert book.active['G2'].data_type=='n'
        with closing(db.get_connection()) as conn:assert '\n'.join(conn.iterdump())==before
        unknown_precision=dict(rows[0],Currency='KWD',**{'Opening balance':Decimal('1.234'),'Closing balance':Decimal('1.234')})
        assert '1.234' in print_document([unknown_precision],Decimal(0),[])
        unknown_book=load_workbook(BytesIO(workbook_bytes([unknown_precision],Decimal(0),[])))
        assert unknown_book.active['G2'].value==1.234 and unknown_book.active['G2'].number_format!='#,##0.00'
        output=Path(os.environ.get('ARETI_QA_REPORT_OUTPUT',str(Path(os.environ['TEMP'])/'report-render')))
        output.mkdir(parents=True,exist_ok=True)
        (output/'sanitized-report.html').write_text(html,encoding='utf-8')
        (output/'sanitized-report.xlsx').write_bytes(workbook_bytes(rows,total,warnings,generated))
        print('PASS SQLite exact text, incomplete/no-import, liabilities, legacy timestamps, green boundaries, safe literal Excel, shared totals and zero writes')
        with closing(db.get_connection()) as conn:
            duplicate_accounts(conn, lambda: snapshot(db), 'SQLite')
    postgres()


def postgres():
    import psycopg2
    options=dict(host='127.0.0.1',port=int(os.environ['ARETI_QA_PG_PORT']),user='qa_local',sslmode='disable')
    database='qa_report_integrity_'+uuid.uuid4().hex
    admin=psycopg2.connect(dbname='postgres',**options);admin.autocommit=True;created=False
    try:
        with admin.cursor() as cursor:
            cursor.execute('SHOW data_directory')
            assert Path(cursor.fetchone()[0]).resolve()==Path(os.environ['ARETI_QA_PG_DATA']).resolve()
            cursor.execute('CREATE DATABASE '+database);created=True
        # Match the current report's existing source/notes/flow columns. This is
        # disposable fixture DDL only; no application/schema change.
        def connect():
            raw=psycopg2.connect(dbname=database,**options)
            raw.set_session(readonly=True)
            return db.PostgresConnection(raw)
        with closing(psycopg2.connect(dbname=database,**options)) as raw, raw:
            with raw.cursor() as cursor:
                cursor.execute('CREATE TABLE account_list(id integer,account_name text,bank text,account_number text,currency text,rate_type text)')
                cursor.execute('CREATE TABLE statement_imports(id integer,statement_hash text,imported_at text,transaction_count integer)')
                cursor.execute("CREATE TABLE statement_balances(statement_hash text,account_name text,bank text,account_number text,currency text,period_start text,period_end text,opening_balance numeric,closing_balance numeric, source text DEFAULT '', notes text DEFAULT '', money_in numeric, money_out numeric)")
                cursor.execute('CREATE TABLE classified_transactions(id integer,statement_hash text,split_parent_id integer,account_name text,bank text,account_number text,currency text)')
                cursor.execute('CREATE TABLE rates(rate_month text,rate_type text,rate_value numeric)')
                cursor.execute("INSERT INTO account_list VALUES(1,'QA','Deposit','QA-1','USD','USD/USD')")
                cursor.execute("INSERT INTO statement_imports VALUES(1,'qa','invalid legacy date',0)")
                cursor.execute("INSERT INTO statement_balances(statement_hash,account_name,bank,account_number,currency,period_start,period_end,opening_balance,closing_balance) VALUES('qa','QA','Deposit','QA-1','USD','2026-09-01','2026-09-30',17.43,17.43)")
        with patch.object(db,'USING_POSTGRES',True),patch.object(db,'get_connection',side_effect=connect):
            exact,total,warnings=snapshot(db)
        assert total==Decimal('17.43') and exact[0]['Verification']=='SOURCE RECONCILIATION NOT VERIFIED'
        assert 'timezone unverified' in exact[0]['Import date'] and any('ordering is unverified' in warning for warning in warnings)
        with closing(psycopg2.connect(dbname=database,**options)) as raw:
            with patch.object(db,'USING_POSTGRES',True),patch.object(db,'get_connection',side_effect=connect):
                duplicate_accounts(db.PostgresConnection(raw), lambda: snapshot(db), 'PostgreSQL')
        for native_type in ('real','double precision'):
            with closing(psycopg2.connect(dbname=database,**options)) as raw, raw:
                with raw.cursor() as cursor:
                    cursor.execute('ALTER TABLE statement_balances ALTER COLUMN opening_balance TYPE '+native_type+', ALTER COLUMN closing_balance TYPE '+native_type)
            with patch.object(db,'USING_POSTGRES',True),patch.object(db,'get_connection',side_effect=connect):
                rows,total,warnings=snapshot(db)
            by_number={row['Account number']:row for row in rows}
            assert total==0 and by_number['QA-1']['Verification']=='UNVERIFIED BINARY STORAGE'
            assert by_number['QA-1']['Closing balance converted to USD']=='NOT AVAILABLE'
            assert warnings
        print('PASS verified local PostgreSQL NUMERIC eligibility, REAL/DOUBLE provenance exclusion, invalid legacy date tolerance and read-only report queries')
    finally:
        if created:
            with admin.cursor() as cursor:cursor.execute('DROP DATABASE '+database)
        admin.close()


def duplicate_accounts(conn, read, backend):
    """One actual identity across two labels must never contribute twice."""
    baseline, base_total, _ = read()
    entries = [
        (100, 'Old label', 'Deposit', 'QA-SAME', 'USD/USD', '2026-10-01T08:00:00+00:00', '10.00'),
        (101, 'New label', 'Deposit', 'QA-SAME', 'EUR/USD', '2026-10-04T08:00:00+00:00', '10.00'),
        (102, 'Safra alias label', 'Safra', 'Current account USD / IBAN CH000000000000000000001', 'USD/USD', '2026-10-01T08:00:00+00:00', '20.00'),
        (103, 'Safra direct label', 'Safra', 'CH000000000000000000001', 'USD/USD', '2026-10-04T08:00:00+00:00', '20.00'),
        (104, 'Safra other actual account', 'Safra', 'CH000000000000000000002', 'USD/USD', '2026-10-04T08:00:00+00:00', '12.00'),
    ]
    cur=conn.cursor()
    try:
        for index,name,bank,number,rate,stamp,amount in entries:
            key='duplicate-qa-'+str(index)
            cur.execute('INSERT INTO account_list(id,account_name,bank,account_number,currency,rate_type) VALUES(?,?,?,?,?,?)',
                        (index,name,bank,number,'USD',rate))
            cur.execute('INSERT INTO statement_imports(id,statement_hash,imported_at,transaction_count) VALUES(?,?,?,0)',(index,key,stamp))
            stored_number='CH000000000000000000001' if index==102 else number
            cur.execute('INSERT INTO statement_balances(statement_hash,account_name,bank,account_number,currency,period_start,period_end,opening_balance,closing_balance) VALUES(?,?,?,?,?,?,?,?,?)',
                        (key,name,bank,stored_number,'USD','2026-09-01','2026-09-30',amount,amount))
        conn.commit()
        rows,total,warnings=read()
        assert len(rows)==len(baseline)+3 and total==base_total+Decimal('12.00')
        dup=[r for r in rows if r['Status']=='AMBIGUOUS SETUP IDENTITY']
        assert len(dup)==2 and all(r['Closing balance converted to USD']=='NOT AVAILABLE' for r in dup)
        ordinary=next(r for r in dup if r['Account number']=='QA-SAME')
        safra=next(r for r in dup if r['Bank']=='Safra')
        assert ordinary['Closing balance']==Decimal('10.00') and ordinary['Import date'].startswith('2026-10-04')
        assert ordinary['Account name']=='Old label (USD/USD); New label (EUR/USD)'
        assert safra['Import date'].startswith('2026-10-04') and 'Safra alias label' in safra['Account name'] and 'Safra direct label' in safra['Account name']
        assert safra['Account number']=='Current account USD / IBAN CH000000000000000000001'
        assert sum('AMBIGUOUS SETUP IDENTITY' in warning for warning in warnings)==2
        assert next(r for r in rows if r['Account number']=='CH000000000000000000002')['Closing balance converted to USD']==Decimal('12.00')
        html=print_document(rows,total,warnings)
        assert html.count('QA-SAME')==2 and 'Old label (USD/USD); New label (EUR/USD)' in html
        book=load_workbook(BytesIO(workbook_bytes(rows,total,warnings)))
        assert sum(row[0].value=='QA-SAME' for row in book.active.iter_rows(min_row=2))==1
        assert book['Verification']['B4'].value==float(total)
        for bad_start in ('malformed','2026-10-01'):
            cur.execute('UPDATE statement_balances SET period_start=? WHERE statement_hash=?',(bad_start,'duplicate-qa-104'));conn.commit()
            changed, changed_total, _ = read()
            other=next(r for r in changed if r['Account number']=='CH000000000000000000002')
            assert other['Status']=='INCOMPLETE' and changed_total==base_total
        cur.execute('UPDATE statement_balances SET period_start=? WHERE statement_hash=?',('2026-09-01','duplicate-qa-104'));conn.commit()
        print('PASS '+backend+' actual-account identity collapse, ambiguous label/rate exclusion, newest across aliases, Safra exact alias normalization, distinct same-currency identities, malformed/reversed dates and shared exports')
    finally:
        cur.close()


if __name__=='__main__':main()
