"""Read-only report date provenance; disposable SQLite fixtures only."""
from contextlib import closing
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch, Mock
import pandas as pd
import json
from openpyxl import load_workbook
import db
from balance_only import PREFIX
from import_history import cyprus_time
from latest_import_balances import snapshot, print_document, workbook_bytes, COLUMNS
from latest_balances_compact import compact_html


class DateTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(dir=os.environ['ARETI_ISOLATION_ROOT'])
        self.patch=patch.object(db,'DB_PATH',str(Path(self.tmp.name)/'report.sqlite'))
        self.patch.start()
        self.assertFalse(db.USING_POSTGRES)
        db.init_db()
        with closing(db.get_connection()) as c:
            ddl=c.execute("SELECT sql FROM sqlite_master WHERE name='statement_balances'").fetchone()[0]
            c.execute('DROP TABLE statement_balances')
            c.execute(ddl.replace('opening_balance REAL','opening_balance TEXT').replace('closing_balance REAL','closing_balance TEXT'))
            c.execute("INSERT INTO account_list(account_name,bank,account_number,currency,rate_type) VALUES('Synthetic','Bank A','A1','USD','USD/USD')")
            c.commit()

    def tearDown(self):
        self.patch.stop()
        self.tmp.cleanup()

    def seed(self,key,stamp,transactions=False,number='A1',bank='Bank A',currency='USD',opening='10.00',closing_balance=None,notes=''):
        with closing(db.get_connection()) as c:
            c.execute('INSERT INTO statement_imports(statement_hash,imported_at,transaction_count) VALUES(?,?,?)',(key,stamp,int(transactions)))
            if transactions:
                c.execute('INSERT INTO classified_transactions(statement_hash,row_hash,txn_date,amount,bank,account_name,account_number,currency,created_at) VALUES(?,?,?,?,?,?,?,?,?)',
                          (key,key+'-row','2026-09-15','-1.00',bank,'Synthetic',number,currency,stamp))
            c.execute('INSERT INTO statement_balances(statement_hash,account_name,bank,account_number,currency,period_start,period_end,opening_balance,closing_balance,money_in,money_out,notes,source,imported_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
                      (key,'Synthetic',bank,number,currency,'2026-09-01','2026-09-30',opening,closing_balance or ('9.00' if transactions else opening),'0','1' if transactions else '0',notes,'Synthetic summary',stamp,stamp))
            c.commit()

    def dump(self):
        with closing(db.get_connection()) as c:return '\n'.join(c.iterdump())

    def report(self):
        before=self.dump()
        result=snapshot(db)
        self.assertEqual(before,self.dump(),'Report must not write any data')
        return result

    def row(self):
        return next(r for r in self.report()[0] if r['Account number']=='A1')

    def test_a_transaction_import_only(self):
        self.seed('transactions','2026-10-01T08:00:00+00:00',True)
        self.assertEqual(self.row()['Transaction Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))

    def test_b_later_balance_only_preserves_transaction_stamp_and_latest_balance(self):
        self.seed('transactions','2026-10-01T08:00:00+00:00',True)
        before=self.row()
        self.seed('balances-only','2026-10-10T08:00:00+00:00',False,opening='20.00')
        after=self.row()
        self.assertEqual(after['Transaction Import Date'],before['Transaction Import Date'])
        self.assertEqual(after['Closing balance'],Decimal('20.00'))
        self.assertEqual(after['Statement end date'],'2026-09-30')

    def test_c_balance_only_and_no_import_use_dash(self):
        self.assertEqual(self.row()['Transaction Import Date'],'—')
        self.seed('balances-only','2026-10-10T08:00:00+00:00')
        self.assertEqual(self.row()['Transaction Import Date'],'—')

    def test_d_same_statement_later_balance_update_uses_original_transaction_history(self):
        self.seed('history','2026-10-01T08:00:00+00:00',True)
        with closing(db.get_connection()) as c:
            c.execute("UPDATE statement_balances SET imported_at='2026-10-10T08:00:00+00:00',updated_at='2026-10-11T08:00:00+00:00' WHERE statement_hash='history'")
            c.commit()
        self.assertEqual(self.row()['Transaction Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))

    def test_e_unrelated_bank_and_currency_history_not_used(self):
        self.seed('balance','2026-10-01T08:00:00+00:00')
        self.seed('other-bank','2026-10-11T08:00:00+00:00',True,bank='Bank B')
        self.seed('other-currency','2026-10-12T08:00:00+00:00',True,currency='EUR')
        self.assertEqual(self.row()['Transaction Import Date'],'—')

    def test_latest_actual_transaction_history_selected_independently(self):
        self.seed('first','2026-10-01T08:00:00+00:00',True)
        self.seed('second','2026-10-03T08:00:00+00:00',True)
        self.seed('balances','2026-10-10T08:00:00+00:00',opening='30.00')
        self.assertEqual(self.row()['Transaction Import Date'],cyprus_time('2026-10-03T08:00:00+00:00'))
        self.assertEqual(self.row()['Closing balance'],Decimal('30.00'))

    def test_transaction_count_without_matching_rows_is_not_transaction_import(self):
        self.seed('incomplete','2026-10-10T08:00:00+00:00')
        with closing(db.get_connection()) as c:
            c.execute("UPDATE statement_imports SET transaction_count=1 WHERE statement_hash='incomplete'")
            c.commit()
        self.assertEqual(self.row()['Transaction Import Date'],'—')

    def test_transaction_import_without_balance_metadata_still_has_its_own_date(self):
        self.seed('transaction-history','2026-10-01T08:00:00+00:00',True)
        with closing(db.get_connection()) as c:
            c.execute("DELETE FROM statement_balances WHERE statement_hash='transaction-history'")
            c.commit()
        self.assertEqual(self.row()['Transaction Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))
        self.assertEqual(self.row()['Status'],'INCOMPLETE')

    def test_missing_transaction_timestamp_never_uses_later_balance_timestamp(self):
        self.seed('transaction-history',None,True)
        self.seed('balance-history','2026-10-11T08:00:00+00:00')
        self.assertEqual(self.row()['Transaction Import Date'],'—')

    def test_postgresql_generated_query_is_readonly_and_parses(self):
        from pglast import parse_sql
        queries=[]
        def read(query,*args,**kwargs):
            queries.append(query)
            return pd.DataFrame()
        with patch.object(db,'USING_POSTGRES',True), patch.object(db,'get_connection',return_value=Mock()), patch.object(db,'_load_rate_lookup',return_value={}), patch('latest_import_balances.pd.read_sql_query',side_effect=read):
            snapshot(db)
        self.assertEqual(len(queries),1)
        parsed=parse_sql(queries[0])
        self.assertEqual(len(parsed),1)
        self.assertEqual(type(parsed[0].stmt).__name__,'SelectStmt')
        self.assertIn('has_transaction_rows',queries[0])

    def test_verified_balance_only_references_never_refresh_transaction_date(self):
        self.seed('original','2026-10-01T08:00:00+00:00',True)
        with closing(db.get_connection()) as c:
            tx_id=c.execute("SELECT id FROM classified_transactions WHERE statement_hash='original'").fetchone()[0]
        notes=PREFIX+json.dumps({'references':[[tx_id,'2026-09-15','-1.00']]})
        self.seed('verified-balances','2026-10-11T08:00:00+00:00',opening='10.00',closing_balance='9.00',notes=notes)
        with closing(db.get_connection()) as c:
            c.execute("UPDATE statement_balances SET money_out='1.00' WHERE statement_hash='verified-balances'")
            c.commit()
        row=self.row()
        self.assertEqual(row['Status'],'IMPORTED')
        self.assertEqual(row['Transaction Import Date'],cyprus_time('2026-10-01T08:00:00+00:00'))

    def test_exports_compact_header_black_date_and_financials_preserved(self):
        self.seed('transactions','2026-10-01T08:00:00+00:00',True)
        self.seed('balances','2026-10-10T08:00:00+00:00',opening='20.00')
        rows,total,warnings=self.report()
        generated=datetime.fromisoformat('2026-10-12T12:00:00+03:00')
        html=print_document(rows,total,warnings,generated)
        compact=compact_html(rows,generated)
        book=load_workbook(BytesIO(workbook_bytes(rows,total,warnings,generated)))
        self.assertIn('Transaction Import Date',COLUMNS)
        self.assertNotIn('Import date',COLUMNS)
        self.assertIn('<th>Transaction Import Date</th>',html)
        self.assertIn('>Transaction Import Date</th>',compact)
        self.assertEqual(book.active.cell(1,4).value,'Transaction Import Date')
        self.assertEqual(book.active.cell(2,4).value,rows[0]['Transaction Import Date'])
        self.assertEqual(book.active.cell(2,4).font.color.rgb,'00000000')
        self.assertIn('color:#000000',compact)
        self.assertIn('color:#146b36',compact)
        self.assertEqual(total,Decimal('20.00'))
        self.assertEqual(rows[0]['Closing balance converted to USD'],Decimal('20.00'))
        self.assertEqual(set(rows[0]),set(COLUMNS),'No private ordering column may leak')


if __name__=='__main__':unittest.main(verbosity=2)
