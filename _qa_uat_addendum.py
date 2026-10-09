"""Bounded source/preview and read-only display UAT addendum checks."""
from contextlib import closing
from collections import Counter
from datetime import datetime,timedelta
from decimal import Decimal
from io import BytesIO
import os,json,unittest
from pathlib import Path
from unittest.mock import patch
from hashlib import sha256
import parsing,db
from fifth_third_import import validate_preview,FifthThirdParseError
from import_history import commit_statement
from latest_balances_compact import compact_html,compact_model
from latest_import_balances import fresh_closing,print_document,workbook_bytes,snapshot
from _qa_post_import_message import CompletionTests
from _qa_fifth_report_match import MatchTests
from zoneinfo import ZoneInfo
ROOT=Path(os.environ['ARETI_ISOLATION_ROOT'])
FACT=json.loads((ROOT/'source-expectations-private.json').read_text(encoding='utf-8'))['Woking Way 2.pdf']

class FourthStatementTests(CompletionTests):
    def source(self):
        content=(ROOT/'Woking Way 2 (1).pdf').read_bytes()
        self.assertEqual(sha256(content).digest(),sha256((ROOT/'Woking Way 2.pdf').read_bytes()).digest())
        frame=parsing.parse_pdf(BytesIO(content));meta=frame.attrs['statement_balance']
        self.assertEqual(frame.Amount.tolist(),list(map(Decimal,FACT['amounts'])))
        for key in ('account_number','currency','period_start','period_end'):
            self.assertEqual(meta[key],FACT[key])
        for key in ('opening_balance','money_in','money_out','closing_balance'):
            self.assertEqual(meta[key],Decimal(FACT[key]))
        account=dict(account_name='Synthetic source owner',bank='Fifth Third (ex-Comerica)',account_number=meta['account_number'],currency='USD',rate_type='USD/USD')
        return content,db.apply_account_and_rates(frame,account),meta,account

    def test_exact_source_preview_and_guards(self):
        content,frame,meta,account=self.source()
        validate_preview(frame,meta,account)
        for override in ({'bank':'Unrelated bank'},{'bank':'Fifth Third (ex-Unrelated)'},{'account_number':'0000'},{'currency':'EUR'}):
            with self.assertRaises(FifthThirdParseError):validate_preview(frame,meta,dict(account,**override))
        changed=frame.copy();changed.loc[0,'Amount']+=Decimal(1)
        with self.assertRaises(FifthThirdParseError):validate_preview(changed,meta,account)

    def test_exact_source_ui_pending_history_balances_duplicate(self):
        content,frame,meta,account=self.source()
        with closing(db.get_connection()) as c:
            c.execute('INSERT INTO account_list(account_name,bank,account_number,currency,rate_type) VALUES(?,?,?,?,?)',tuple(account.values()));c.commit()
        self.file=BytesIO(content);self.file.name='Woking Way 2 (1).pdf'
        app=self.run_app(self.app());self.import_file(app);self.complete(app,len(frame))
        self.assertEqual(len(db.get_pending_transactions()),len(frame))
        # Pending is date-sorted; compare exact date/amount multisets, not display order or numeric spelling.
        pending=db.get_pending_transactions()
        self.assertEqual(Counter((str(r.txn_date),Decimal(str(r.amount))) for r in pending.itertuples()),Counter((str(r.Date),Decimal(str(r.Amount))) for r in frame.itertuples()))
        history=db.get_import_history();self.assertEqual(int(history.iloc[0].transaction_count),len(frame))
        rows,total,_=snapshot(db);row=next(r for r in rows if r['Account number']==account['account_number'])
        self.assertEqual(row['Closing balance'],meta['closing_balance'])
        before=self.snapshot();digest=sha256(content).hexdigest()
        self.assertEqual(commit_statement(db,frame,self.file.name,digest,meta,account),(0,True,0))
        self.assertEqual(self.snapshot(),before)

class USDExclusionTests(MatchTests):
    def test_card_balance_retains_native_and_excludes_compact_usd(self):
        self.seed('Sapphire Chase')
        with closing(db.get_connection()) as c:
            c.execute("UPDATE statement_balances SET bank='Sapphire Chase'")
            c.execute("UPDATE classified_transactions SET bank='Sapphire Chase'");c.commit()
        before=self.snapshot();rows,total,_=snapshot(db)
        row=next(r for r in rows if r['Account number']=='1892774128')
        self.assertEqual(row['Verification'],'LIABILITY CONVENTION UNVERIFIED')
        visible,compact_total,_,_=compact_model([row])
        self.assertEqual(visible[0]['closing'],Decimal('90.05'));self.assertIsNone(visible[0]['usd'])
        self.assertEqual(total,Decimal(0));self.assertEqual(compact_total,Decimal(0));self.assertEqual(self.snapshot(),before)
    def test_no_rate_is_not_invented_and_precision_guard_not_relaxed(self):
        self.seed()
        with closing(db.get_connection()) as c:
            c.execute("UPDATE account_list SET currency='EUR',rate_type='EUR/USD' WHERE account_number='1892774128'")
            c.execute("UPDATE statement_balances SET currency='EUR'")
            c.execute("UPDATE classified_transactions SET currency='EUR'");c.commit()
        before=self.snapshot();rows,total,_=snapshot(db)
        row=next(r for r in rows if r['Account number']=='1892774128')
        self.assertEqual(row['Status'],'IMPORTED');self.assertEqual(row['Applied FX'],'')
        self.assertEqual(row['Closing balance converted to USD'],'NOT AVAILABLE');self.assertEqual(total,Decimal(0))
        self.assertEqual(self.snapshot(),before)
        with closing(db.get_connection()) as c:
            c.execute("UPDATE statement_balances SET closing_balance='962.5999755859375'");c.commit()
        rows,total,_=snapshot(db);row=next(r for r in rows if r['Account number']=='1892774128')
        self.assertEqual(row['Verification'],'UNVERIFIED STORED PRECISION');self.assertEqual(total,Decimal(0))
        self.assertIn('962.60',compact_html([row]));self.assertEqual(row['Closing balance'],Decimal('962.5999755859375'))

class PresentationTests(unittest.TestCase):
    def rows(self):
        base={'Status':'IMPORTED','Verification':'SOURCE RECONCILIATION NOT VERIFIED','Bank':'Bank A','Account name':'Owner','Account number':'0001','Import date':'2026-10-09','Statement end date':'2026-09-30','Currency':'USD','Closing balance':Decimal('962.5999755859375'),'Closing balance converted to USD':Decimal('962.5999755859375')}
        return [base,dict(base,**{'Account name':'AAA','Statement end date':'2026-09-30','Closing balance':Decimal('2.10'),'Closing balance converted to USD':Decimal('2.10')}),dict(base,**{'Account name':'AAA','Statement end date':'2026-09-01','Closing balance':Decimal('3.20'),'Closing balance converted to USD':Decimal('3.20')})]
    def test_columns_compact_width_two_decimals_sort_and_exact_total(self):
        rows=self.rows();before=repr(rows);html=compact_html(rows)
        labels=('Bank','Account','Account Number','Import Date','Statement Closing Date','Currency','Closing Balance','Closing Balance USD')
        self.assertEqual([html.index('>'+label+'</th>') for label in labels],sorted(html.index('>'+label+'</th>') for label in labels))
        self.assertIn('table-layout:auto',html);self.assertIn('width:max-content',html);self.assertIn('white-space:nowrap',html)
        self.assertNotIn('962.5999755859375',html);self.assertIn('962.60',html);self.assertIn('2.10',html)
        visible,total,*_=compact_model(rows)
        self.assertEqual([(r['account'],r['closing_date']) for r in visible],[('AAA','2026-09-01'),('AAA','2026-09-30'),('Owner','2026-09-30')])
        self.assertEqual(total,Decimal('967.8999755859375'));self.assertEqual(repr(rows),before)
        self.assertEqual(html.count('<tr>'),5)
    def test_inclusive_30_day_closing_freshness_all_displays(self):
        now=datetime(2026,10,9,tzinfo=ZoneInfo('Europe/Nicosia'))
        from openpyxl import load_workbook
        from latest_import_balances import COLUMNS
        for age in (-1,0,29,30,31,40):
            close=(now.date()-timedelta(days=age)).isoformat()
            self.assertEqual(fresh_closing(close,now),0<=age<=30)
            row={k:'' for k in COLUMNS};row.update(self.rows()[0]);row['Statement end date']=close
            html=compact_html([row],now)
            self.assertIn('color:'+('#146b36' if 0<=age<=30 else '#b42318'),html)
            self.assertIn('color:#000000',html)
            doc=print_document([row],Decimal(0),[],now)
            self.assertIn('class="'+('fresh' if 0<=age<=30 else 'stale')+'">'+close,doc)
            sheet=load_workbook(BytesIO(workbook_bytes([row],Decimal(0),[],now))).active
            self.assertEqual(sheet.cell(2,6).font.color.rgb,'00146B36' if 0<=age<=30 else '00B42318')
            self.assertEqual(sheet.cell(2,4).font.color.rgb,'00000000')
        for value in ('',None,'invalid'):self.assertFalse(fresh_closing(value,now))


def load_tests(loader,tests,pattern):
    return unittest.TestSuite([FourthStatementTests('test_exact_source_preview_and_guards'),FourthStatementTests('test_exact_source_ui_pending_history_balances_duplicate'),loader.loadTestsFromTestCase(PresentationTests),USDExclusionTests('test_card_balance_retains_native_and_excludes_compact_usd'),USDExclusionTests('test_no_rate_is_not_invented_and_precision_guard_not_relaxed')])
if __name__=='__main__':unittest.main()
